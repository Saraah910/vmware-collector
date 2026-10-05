import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def format_value(value):

    if value is None:
        return ""

    if isinstance(value, bool):
        return "Yes" if value else "No"

    if isinstance(value, (str, int, float)):
        return str(value)

    if isinstance(value, dict):
        return "\n".join(
            f"{k.replace('_', ' ').title()}: {format_value(v)}"
            for k, v in value.items()
        )

    if isinstance(value, list):

        if not value:
            return ""

        formatted = []

        for item in value:

            if isinstance(item, dict):

                formatted.append(
                    "\n".join(
                        f"{k.replace('_', ' ').title()}: {format_value(v)}"
                        for k, v in item.items()
                    )
                )

            else:
                formatted.append(str(item))

        return "\n\n".join(formatted)

    return str(value)


def generate_vm_excel(json_file, output_excel="VM_Report.xlsx"):

    with open(json_file, "r", encoding="utf-8") as f:
        data = json.load(f)

    wb = Workbook()
    ws = wb.active
    ws.title = "Virtual Machines"

    # Build column list dynamically
    headers = [
        "vm_name",
        "vcentre_ip",
        "configured_os",
        "guest_os",
        "cpus",
        "cores",
        "memory",
        # "cpu_hot_add_enabled",
        # "memory_hot_add_enabled",
        "vmware_tools",
        "secure_boot_status",
        "snapshots",
        "cd_dvd_device",
        "networks",
        "virtual_disks",
        "independant_disks"
    ]

    header_fill = PatternFill(
        start_color="1F4E78",
        end_color="1F4E78",
        fill_type="solid"
    )
    for vm in data.values():
        for key in vm.keys():
            if key not in headers:
                headers.append(key)

    # Header row
    for col_num, header in enumerate(headers, start=1):

        cell = ws.cell(
            row=1,
            column=col_num,
            value=header.replace("_", " ").title()
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = header_fill


    # Data rows
    for row_num, vm in enumerate(data.values(), start=2):

        for col_num, header in enumerate(headers, start=1):

            value = vm.get(header)

            cell = ws.cell(
                row=row_num,
                column=col_num,
                value=format_value(value)
            )

            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top"
            )

    # Autosize columns
    for column in ws.columns:

        max_length = 0

        for cell in column:

            try:
                if cell.value:
                    longest_line = max(
                        len(line)
                        for line in str(cell.value).split("\n")
                    )

                    max_length = max(
                        max_length,
                        longest_line
                    )

            except Exception:
                pass

        adjusted_width = min(max_length + 3, 60)

        ws.column_dimensions[
            column[0].column_letter
        ].width = adjusted_width

    # Enable filter row
    ws.auto_filter.ref = ws.dimensions

    wb.save(output_excel)

    print(f"Excel report generated: {output_excel}")


if __name__ == "__main__":

    if len(sys.argv) < 2:
        print("Usage: python report.py <jsonfile>")
        sys.exit(1)

    generate_vm_excel(
        sys.argv[1],
        "VM_Report.xlsx"
    )